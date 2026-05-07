"""Builds the Vialitycs MX pitch deck (English).

Output: reports/presentacion_vialitycs_mx.pptx

Visual tokens (inspired by reports/reference.pptx):
- Black background (#000000)
- Lime accent (#C1FF72)
- Display: RQND Pro Wide
- Body: Fira Code
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


# ---------- Palette ----------
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIME = RGBColor(0xC1, 0xFF, 0x72)
DARK_SURFACE = RGBColor(0x12, 0x12, 0x12)
MID_SURFACE = RGBColor(0x1C, 0x1C, 0x1C)
BORDER = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x9A, 0x9A, 0x9A)

# ---------- Typography ----------
DISPLAY_FONT = "RQND Pro Wide"
BODY_FONT = "Fira Code"

# ---------- Geometry ----------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_no_line(shape) -> None:
    shape.line.fill.background()


def set_line(shape, color: RGBColor, width_pt: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def set_slide_bg(slide, color: RGBColor = BLACK) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    font_size: int = 14,
    bold: bool = False,
    color: RGBColor = WHITE,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font_name: str = BODY_FONT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    items: list[str],
    *,
    font_size: int = 14,
    color: RGBColor = WHITE,
    line_spacing: float = 1.4,
    bullet_color: RGBColor = LIME,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(8)
        bullet_run = p.add_run()
        bullet_run.text = "▍ "
        bullet_run.font.name = BODY_FONT
        bullet_run.font.size = Pt(font_size)
        bullet_run.font.color.rgb = bullet_color
        bullet_run.font.bold = True
        text_run = p.add_run()
        text_run.text = item
        text_run.font.name = BODY_FONT
        text_run.font.size = Pt(font_size)
        text_run.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, *, fill=DARK_SURFACE, line=None, line_width=1.0):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    set_fill(rect, fill)
    if line is None:
        set_no_line(rect)
    else:
        set_line(rect, line, line_width)
    rect.shadow.inherit = False
    return rect


def add_brand_strip(slide):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.35),
        Inches(4),
        Inches(0.3),
        "VIALITYCS MX",
        font_size=10,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(2.0), Inches(0.45), Inches(0.1), Inches(0.1)
    )
    set_fill(dot, LIME)
    set_no_line(dot)


def add_page_footer(slide, page: int, total: int):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.18),
        Inches(8),
        Inches(0.25),
        "DATA PROCESSING · OMAR CHI · VICTOR FIGUEREDO",
        font_size=8,
        color=MUTED,
    )
    add_textbox(
        slide,
        Inches(11.5),
        Inches(7.18),
        Inches(1.4),
        Inches(0.25),
        f"{page:02d} / {total:02d}",
        font_size=8,
        color=LIME,
        align=PP_ALIGN.RIGHT,
        bold=True,
    )


def add_slide_title(slide, title: str, kicker: str | None = None, *, title_top: float = 1.05):
    if kicker:
        add_textbox(
            slide,
            Inches(0.5),
            Inches(0.85),
            Inches(12),
            Inches(0.3),
            kicker.upper(),
            font_size=10,
            bold=True,
            color=LIME,
        )
        title_top_in = title_top + 0.18
    else:
        title_top_in = title_top
    add_textbox(
        slide,
        Inches(0.5),
        Inches(title_top_in),
        Inches(12.3),
        Inches(1.1),
        title,
        font_size=34,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    add_rect(
        slide,
        Inches(0.5),
        Inches(title_top_in + 1.0),
        Inches(0.5),
        Inches(0.04),
        fill=LIME,
    )


def add_placeholder(
    slide, left, top, width, height, label: str, *, sublabel: str | None = None
):
    rect = add_rect(
        slide,
        left,
        top,
        width,
        height,
        fill=DARK_SURFACE,
        line=LIME,
        line_width=1.25,
    )
    rect.line.dash_style = 7
    label_box = slide.shapes.add_textbox(left, top, width, height)
    tf = label_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = DISPLAY_FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = LIME
    if sublabel:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(6)
        r2 = p2.add_run()
        r2.text = sublabel
        r2.font.name = BODY_FONT
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED
    return rect


# ---------- Slides ----------

def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    add_rect(slide, Inches(6.4), Inches(0.7), Inches(0.04), Inches(6.1), fill=LIME)

    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.7),
        Inches(5),
        Inches(0.3),
        "PITCH · MAY 2026",
        font_size=10,
        bold=True,
        color=LIME,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.2),
        Inches(6),
        Inches(2.0),
        "Vialitycs\nMX.",
        font_size=72,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.6),
        Inches(5.7),
        Inches(1.0),
        "Actionable intelligence on Mexico's\npublic road traffic and vehicle data.",
        font_size=14,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(6.7),
        Inches(5.7),
        Inches(0.3),
        "DATA PROCESSING",
        font_size=10,
        color=MUTED,
        bold=True,
    )

    add_placeholder(
        slide,
        Inches(7.0),
        Inches(0.7),
        Inches(5.85),
        Inches(4.2),
        "[ HERO IMAGE / LOGO ]",
        sublabel="Suggested: traffic render, crash photo or MX map. 1600 × 900 px.",
    )
    add_textbox(
        slide,
        Inches(7.0),
        Inches(5.2),
        Inches(5.85),
        Inches(0.3),
        "PRESENTED BY",
        font_size=10,
        bold=True,
        color=LIME,
    )
    add_textbox(
        slide,
        Inches(7.0),
        Inches(5.55),
        Inches(5.85),
        Inches(0.6),
        "Omar Antonio Chi Bautista",
        font_size=18,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    add_textbox(
        slide,
        Inches(7.0),
        Inches(6.05),
        Inches(5.85),
        Inches(0.6),
        "Victor Figueredo",
        font_size=18,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )


def slide_problem(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "The data exists.\nNobody can use it.", kicker="01 · The problem")

    add_placeholder(
        slide,
        Inches(0.5),
        Inches(3.4),
        Inches(5.5),
        Inches(3.4),
        "[ CRASH / TRAFFIC IMAGE ]",
        sublabel="To liven things up. Suggested: editorial photo, high contrast, B&W.",
    )

    stats = [
        ("28", "years of ATUS data left unused (1997–2024)"),
        ("32", "states · thousands of municipalities"),
        ("23/28", "files with broken year column (YEAR=1)"),
    ]
    top_y = 3.4
    for i, (big, small) in enumerate(stats):
        y = Inches(top_y + i * 1.2)
        add_rect(slide, Inches(6.3), y, Inches(6.5), Inches(1.05), fill=DARK_SURFACE)
        add_rect(slide, Inches(6.3), y, Inches(0.08), Inches(1.05), fill=LIME)
        add_textbox(
            slide,
            Inches(6.55),
            y + Inches(0.18),
            Inches(2.3),
            Inches(0.7),
            big,
            font_size=36,
            bold=True,
            color=LIME,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            Inches(8.7),
            y + Inches(0.32),
            Inches(4.0),
            Inches(0.6),
            small,
            font_size=11,
            color=WHITE,
        )

    add_page_footer(slide, page, total)


def slide_audience(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Who needs these answers?", kicker="02 · Who feels the pain")

    cards = [
        ("Insurers", "Size exposure and price policies by state and municipality."),
        ("Local governments", "Prioritize road-safety investment with quantitative evidence."),
        ("Journalists & NGOs", "Ground stories in official figures in minutes, not days."),
        ("Mobility teams", "Cross sales and fleet trends with crash data."),
    ]
    card_w = Inches(2.95)
    card_h = Inches(3.5)
    gap = Inches(0.15)
    start_x = Inches(0.5)
    top = Inches(2.7)
    for i, (titulo, desc) in enumerate(cards):
        x = start_x + (card_w + gap) * i
        add_rect(slide, x, top, card_w, card_h, fill=DARK_SURFACE)
        add_rect(slide, x, top, card_w, Inches(0.08), fill=LIME)
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(0.4),
            Inches(2.5),
            Inches(0.5),
            f"0{i+1}",
            font_size=28,
            bold=True,
            color=LIME,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(1.1),
            card_w - Inches(0.6),
            Inches(0.5),
            titulo,
            font_size=16,
            bold=True,
            color=WHITE,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(1.85),
            card_w - Inches(0.6),
            Inches(1.5),
            desc,
            font_size=11,
            color=WHITE,
        )

    add_page_footer(slide, page, total)


def slide_solution(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(
        slide,
        "Ask in plain Spanish.\nGet an executive analysis.",
        kicker="03 · The solution",
    )

    left_x = Inches(0.5)
    top = Inches(3.0)
    add_rect(slide, left_x, top, Inches(6.2), Inches(0.85), fill=DARK_SURFACE)
    add_textbox(
        slide,
        left_x + Inches(0.25),
        top + Inches(0.08),
        Inches(5.7),
        Inches(0.3),
        "INPUT",
        font_size=9,
        bold=True,
        color=LIME,
    )
    add_textbox(
        slide,
        left_x + Inches(0.25),
        top + Inches(0.38),
        Inches(5.7),
        Inches(0.5),
        "> ¿Qué estados tuvieron más accidentes en 2023?",
        font_size=12,
        color=WHITE,
    )

    resp_top = Inches(4.0)
    add_rect(slide, left_x, resp_top, Inches(6.2), Inches(2.85), fill=MID_SURFACE)
    add_rect(slide, left_x, resp_top, Inches(0.08), Inches(2.85), fill=LIME)
    add_textbox(
        slide,
        left_x + Inches(0.3),
        resp_top + Inches(0.15),
        Inches(5.7),
        Inches(0.3),
        "OUTPUT",
        font_size=9,
        bold=True,
        color=LIME,
    )
    sections = [
        ("Summary", "national 2023 ranking with official figures."),
        ("Data", "top 5 states with real totals."),
        ("Interpretation", "narrative without claiming causality."),
        ("Limitations", "ATUS 2022 methodology break."),
    ]
    y = resp_top + Inches(0.55)
    for titulo, desc in sections:
        add_textbox(
            slide,
            left_x + Inches(0.3),
            y,
            Inches(1.7),
            Inches(0.4),
            titulo,
            font_size=11,
            bold=True,
            color=LIME,
        )
        add_textbox(
            slide,
            left_x + Inches(2.0),
            y,
            Inches(4.0),
            Inches(0.4),
            desc,
            font_size=11,
            color=WHITE,
        )
        y += Inches(0.55)

    add_textbox(
        slide,
        Inches(7.0),
        Inches(2.95),
        Inches(5.85),
        Inches(0.5),
        "What Vialitycs MX delivers",
        font_size=18,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    add_bullets(
        slide,
        Inches(7.0),
        Inches(3.65),
        Inches(5.85),
        Inches(3.5),
        [
            "ATUS coverage 1997–2024 by state and municipality.",
            "VMRC national: vehicle sales and active fleet.",
            "Structured response in 4 fixed sections.",
            "Zero numeric hallucinations — pandas computes, the LLM writes.",
            "Working Streamlit web app, ready to extend to API.",
        ],
        font_size=12,
    )

    add_page_footer(slide, page, total)


def slide_demo(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Vialitycs MX in action", kicker="04 · Demo")

    add_placeholder(
        slide,
        Inches(1.0),
        Inches(2.4),
        Inches(11.3),
        Inches(4.5),
        "[ INSERT DEMO VIDEO HERE ]",
        sublabel="Insert > Video > This device. Auto-play recommended.",
    )

    add_page_footer(slide, page, total)


def slide_trust(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Built to never lie.", kicker="05 · Why trust the numbers")

    pillars = [
        ("01", "Never calculates", "Every metric is computed in pandas on official INEGI data. The model only writes."),
        ("02", "Never invents", "The LLM only cites what arrives in the prompt. No free access to its memory."),
        ("03", "Never claims cause", "It speaks of correlation or association. Claiming “cause” is forbidden by design."),
    ]
    card_w = Inches(3.95)
    card_h = Inches(3.3)
    gap = Inches(0.2)
    start_x = Inches(0.5)
    top = Inches(2.7)
    for i, (num, titulo, desc) in enumerate(pillars):
        x = start_x + (card_w + gap) * i
        add_rect(slide, x, top, card_w, card_h, fill=DARK_SURFACE, line=BORDER, line_width=1.0)
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(0.3),
            Inches(2),
            Inches(0.7),
            num,
            font_size=40,
            bold=True,
            color=LIME,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(1.25),
            card_w - Inches(0.6),
            Inches(0.6),
            titulo,
            font_size=18,
            bold=True,
            color=WHITE,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.3),
            top + Inches(2.05),
            card_w - Inches(0.6),
            Inches(1.2),
            desc,
            font_size=11,
            color=WHITE,
        )

    add_textbox(
        slide,
        Inches(0.5),
        Inches(6.4),
        Inches(12.3),
        Inches(0.5),
        "“ If the data isn't in INEGI, Vialitycs won't say it. ”",
        font_size=15,
        bold=True,
        color=LIME,
        align=PP_ALIGN.CENTER,
        font_name=DISPLAY_FONT,
    )

    add_page_footer(slide, page, total)


def slide_architecture(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "How it works", kicker="06 · Architecture")

    steps = [
        ("User\nquestion", "“Top 5 in\n2023”"),
        ("Intent\ndetection", "intent +\nslots"),
        ("Pandas\ncomputation", "INEGI\ndata"),
        ("Prompt\nbuilding", "data +\nrules"),
        ("Fine-tuned\nLLM", "Qwen-1.5B\n+ LoRA"),
        ("Executive\nanswer", "4 standard\nsections"),
    ]
    n = len(steps)
    total_w = Inches(12.3)
    start_x = Inches(0.5)
    top = Inches(2.9)
    box_w = Inches(1.78)
    box_h = Inches(2.0)
    gap = (total_w - box_w * n) / (n - 1)

    for i, (titulo, sub) in enumerate(steps):
        x = start_x + (box_w + gap) * i
        add_rect(slide, x, top, box_w, box_h, fill=DARK_SURFACE, line=BORDER, line_width=1.0)
        add_rect(slide, x, top, box_w, Inches(0.08), fill=LIME)
        add_textbox(
            slide,
            x + Inches(0.1),
            top + Inches(0.35),
            box_w - Inches(0.2),
            Inches(0.9),
            titulo,
            font_size=12,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.1),
            top + Inches(1.25),
            box_w - Inches(0.2),
            Inches(0.7),
            sub,
            font_size=9,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
        if i < n - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                x + box_w + Inches(0.02),
                top + Inches(0.85),
                gap - Inches(0.04),
                Inches(0.3),
            )
            set_fill(arrow, LIME)
            set_no_line(arrow)

    add_textbox(
        slide,
        Inches(0.5),
        Inches(5.5),
        Inches(12.3),
        Inches(0.5),
        "Numbers come from pandas. The LLM only writes.",
        font_size=15,
        bold=True,
        color=LIME,
        align=PP_ALIGN.CENTER,
        font_name=DISPLAY_FONT,
    )

    add_page_footer(slide, page, total)


def slide_use_cases(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "What can it answer today?", kicker="07 · Use cases")

    questions = [
        "Top 5 states with the most crashes in 2023.",
        "Crash trend in Yucatán between 2015 and 2024.",
        "Compare CDMX vs Jalisco in 2022.",
        "Are vehicle sales linked to crashes?",
    ]
    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.7),
        Inches(5.8),
        Inches(0.5),
        "Supported question types",
        font_size=16,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    y = Inches(3.3)
    for q in questions:
        add_rect(slide, Inches(0.5), y, Inches(5.8), Inches(0.78), fill=DARK_SURFACE)
        add_rect(slide, Inches(0.5), y, Inches(0.06), Inches(0.78), fill=LIME)
        add_textbox(
            slide,
            Inches(0.78),
            y + Inches(0.18),
            Inches(5.4),
            Inches(0.5),
            f"> {q}",
            font_size=11,
            color=WHITE,
        )
        y += Inches(0.9)

    add_placeholder(
        slide,
        Inches(7.0),
        Inches(2.7),
        Inches(5.85),
        Inches(4.2),
        "[ APP SCREENSHOTS ]",
        sublabel="Suggested: 1 prompt screenshot + 1 response screenshot.",
    )

    add_page_footer(slide, page, total)


def slide_metrics(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Fine-tuning nearly doubles ROUGE-1.", kicker="08 · Results")

    headers = ["Metric", "FLAN-T5 base", "Qwen-1.5B vanilla", "Vialitycs (Qwen-LoRA)"]
    # (label, flan, qwen, lora, winner_index_for_highlight or None)
    # winner_index: 1 = flan col, 2 = qwen vanilla col, 3 = lora col, None = no highlight
    rows = [
        ("ROUGE-1",                    "0.17",  "0.38",  "0.71",  3),
        ("ROUGE-2",                    "0.03",  "0.12",  "0.64",  3),
        ("ROUGE-L",                    "0.09",  "0.23",  "0.68",  3),
        ("Sections present (0–1)",     "0.03",  "0.95",  "0.80",  2),
        ("No-causality (correlation)", "1.00",  "1.00",  "1.00",  None),
        ("Avg length (words)",         "124",   "161",   "160",   None),
    ]
    table_left = Inches(0.5)
    table_top = Inches(2.7)
    col_widths = [Inches(3.6), Inches(2.4), Inches(2.8), Inches(3.5)]
    row_h = Inches(0.5)

    # Header row
    x = table_left
    for i, h in enumerate(headers):
        add_rect(slide, x, table_top, col_widths[i], row_h, fill=LIME)
        add_textbox(
            slide,
            x + Inches(0.15),
            table_top + Inches(0.11),
            col_widths[i] - Inches(0.3),
            row_h,
            h,
            font_size=10,
            bold=True,
            color=BLACK,
        )
        x += col_widths[i]

    # Data rows
    for r_idx, row in enumerate(rows):
        label, flan, qwen, lora, winner = row
        cells = [label, flan, qwen, lora]
        x = table_left
        y = table_top + row_h * (r_idx + 1)
        bg = DARK_SURFACE if r_idx % 2 == 0 else MID_SURFACE
        for i, cell in enumerate(cells):
            add_rect(slide, x, y, col_widths[i], row_h, fill=bg)
            is_winner = winner is not None and i == winner
            color = LIME if is_winner else WHITE
            bold = is_winner or i == 0
            add_textbox(
                slide,
                x + Inches(0.15),
                y + Inches(0.11),
                col_widths[i] - Inches(0.3),
                row_h,
                cell,
                font_size=11,
                bold=bold,
                color=color,
            )
            x += col_widths[i]

    add_textbox(
        slide,
        Inches(0.5),
        Inches(6.45),
        Inches(12.3),
        Inches(0.4),
        "Test set: 42 held-out questions. ROUGE measured without stemmer. Higher is better.",
        font_size=9,
        color=MUTED,
    )

    add_page_footer(slide, page, total)


def slide_traction(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Not just an idea — it's already built.", kicker="09 · Traction")

    items = [
        ("ETL pipeline", "28 years of ATUS and VMRC, reproducible in one command."),
        ("Curated dataset", "425 training examples (train · val · test)."),
        ("Own model", "Qwen2.5-1.5B fine-tuned with QLoRA on Colab."),
        ("Working app", "End-to-end Streamlit with the model loaded."),
        ("Documentation", "Stage-by-stage plan and notes to resume / extend."),
        ("Validation", "Independent test set with 6 automated metrics."),
    ]
    cols = 3
    card_w = Inches(4.05)
    card_h = Inches(1.85)
    gap_x = Inches(0.13)
    gap_y = Inches(0.2)
    start_x = Inches(0.5)
    start_y = Inches(2.7)
    for i, (titulo, desc) in enumerate(items):
        col = i % cols
        row = i // cols
        x = start_x + (card_w + gap_x) * col
        y = start_y + (card_h + gap_y) * row
        add_rect(slide, x, y, card_w, card_h, fill=DARK_SURFACE)
        add_rect(slide, x, y, Inches(0.08), card_h, fill=LIME)
        add_textbox(
            slide,
            x + Inches(0.35),
            y + Inches(0.25),
            card_w - Inches(0.5),
            Inches(0.5),
            titulo,
            font_size=14,
            bold=True,
            color=WHITE,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x + Inches(0.35),
            y + Inches(0.85),
            card_w - Inches(0.5),
            Inches(0.9),
            desc,
            font_size=10,
            color=WHITE,
        )

    add_page_footer(slide, page, total)


def slide_roadmap(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Where Vialitycs MX is headed", kicker="10 · Roadmap")

    phases = [
        ("TODAY", "Functional academic MVP with national coverage."),
        ("Q3 2026", "State-level VMRC coverage + DENUE integration."),
        ("Q4 2026", "Export to PDF · REST API · multi-user."),
        ("2027", "Enterprise version with authentication and SLA."),
    ]
    n = len(phases)
    line_y = Inches(4.0)
    add_rect(slide, Inches(0.5), line_y, Inches(12.3), Inches(0.04), fill=LIME)

    box_w = Inches(2.9)
    gap = (Inches(12.3) - box_w * n) / (n - 1)
    for i, (when, desc) in enumerate(phases):
        x = Inches(0.5) + (box_w + gap) * i
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            x + box_w / 2 - Inches(0.14),
            line_y - Inches(0.115),
            Inches(0.28),
            Inches(0.28),
        )
        set_fill(dot, LIME)
        set_no_line(dot)
        add_textbox(
            slide,
            x,
            Inches(2.85),
            box_w,
            Inches(0.5),
            when,
            font_size=18,
            bold=True,
            color=LIME,
            align=PP_ALIGN.CENTER,
            font_name=DISPLAY_FONT,
        )
        add_textbox(
            slide,
            x,
            Inches(4.5),
            box_w,
            Inches(2.0),
            desc,
            font_size=12,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )

    add_page_footer(slide, page, total)


def slide_stack(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_brand_strip(slide)
    add_slide_title(slide, "Open stack, zero cost.", kicker="11 · Technology")

    items_left = [
        "Python · pandas · Streamlit",
        "Hugging Face Transformers · PEFT (LoRA)",
        "Base model Qwen2.5-1.5B-Instruct (Apache 2.0)",
        "Trained on Google Colab T4 (free tier)",
    ]
    items_right = [
        "100% open source — no vendor lock-in.",
        "Local inference on modest hardware (RTX 3060 Ti).",
        "No paid-API dependencies.",
        "Reproducible: pipeline + training in one command.",
    ]
    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.7),
        Inches(5.8),
        Inches(0.5),
        "Components",
        font_size=18,
        bold=True,
        color=LIME,
        font_name=DISPLAY_FONT,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(3.4),
        Inches(5.8),
        Inches(3.5),
        items_left,
        font_size=12,
    )
    add_textbox(
        slide,
        Inches(7.0),
        Inches(2.7),
        Inches(5.85),
        Inches(0.5),
        "Business implications",
        font_size=18,
        bold=True,
        color=LIME,
        font_name=DISPLAY_FONT,
    )
    add_bullets(
        slide,
        Inches(7.0),
        Inches(3.4),
        Inches(5.85),
        Inches(3.5),
        items_right,
        font_size=12,
    )

    add_page_footer(slide, page, total)


def slide_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    add_textbox(
        slide,
        Inches(0.5),
        Inches(0.7),
        Inches(8),
        Inches(0.4),
        "VIALITYCS MX",
        font_size=10,
        bold=True,
        color=LIME,
        font_name=DISPLAY_FONT,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(2.0),
        Inches(12.3),
        Inches(2.0),
        "Public data.\nPrivate decisions.",
        font_size=64,
        bold=True,
        color=WHITE,
        font_name=DISPLAY_FONT,
    )
    add_rect(slide, Inches(0.5), Inches(4.6), Inches(0.7), Inches(0.06), fill=LIME)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(4.85),
        Inches(8),
        Inches(0.6),
        "Thank you.",
        font_size=28,
        bold=True,
        color=LIME,
        font_name=DISPLAY_FONT,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(5.55),
        Inches(8),
        Inches(0.5),
        "Omar Antonio Chi Bautista  ·  Victor Figueredo",
        font_size=14,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.5),
        Inches(5.95),
        Inches(8),
        Inches(0.5),
        "Data Processing · May 2026",
        font_size=11,
        color=MUTED,
    )

    add_placeholder(
        slide,
        Inches(9.3),
        Inches(2.0),
        Inches(3.55),
        Inches(3.55),
        "[ QR / CONTACT ]",
        sublabel="Insert photo, QR code or repository link.",
    )


# ---------- Build ----------

def build(output: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    builders = [
        slide_cover,
        slide_problem,
        slide_audience,
        slide_solution,
        slide_demo,
        slide_trust,
        slide_architecture,
        slide_use_cases,
        slide_metrics,
        slide_traction,
        slide_roadmap,
        slide_stack,
        slide_close,
    ]
    total = len(builders)
    for idx, builder in enumerate(builders, start=1):
        if builder is slide_cover or builder is slide_close:
            builder(prs)
        else:
            builder(prs, idx, total)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


if __name__ == "__main__":
    out = Path(__file__).resolve().parents[1] / "reports" / "presentacion_vialitycs_mx.pptx"
    saved = build(out)
    print(f"OK: {saved}")
